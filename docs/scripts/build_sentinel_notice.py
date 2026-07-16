from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = r"C:\Users\Morpheus\site vitrine DevOra\docs\DEVORA_SENTINEL_NOTICE_PROCEDURE_CLIENTS.docx"

NAVY = RGBColor(6, 26, 53)
BLUE = RGBColor(37, 99, 235)
CYAN_DARK = RGBColor(8, 145, 178)
VIOLET = RGBColor(124, 58, 237)
MUTED = RGBColor(71, 85, 105)
LIGHT_BLUE = "EAF5FF"
LIGHT_CYAN = "ECFEFF"
LIGHT_GRAY = "F8FAFC"
BORDER = "CBD5E1"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if table.rows:
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:tblHeader")) is None:
            tbl_header = OxmlElement("w:tblHeader")
            tbl_header.set(qn("w:val"), "true")
            tr_pr.append(tbl_header)


def set_font(run, size=None, color=None, bold=None, italic=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def style_doc(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    for attr in ("top_margin", "right_margin", "bottom_margin", "left_margin"):
        setattr(section, attr, Inches(1))
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = NAVY
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, NAVY, 10, 5),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def add_p(doc, text="", style=None, bold=False, color=NAVY, size=11, after=None):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    set_font(r, size=size, color=color, bold=bold)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    set_font(r, size={1: 16, 2: 13, 3: 12}[level], color=BLUE if level < 3 else NAVY, bold=True)
    return p


def add_bullet(doc, text):
    return add_p(doc, text, style="List Bullet", color=NAVY)


def add_number(doc, text):
    return add_p(doc, text, style="List Number", color=NAVY)


def add_callout(doc, title, body, fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_font(r, size=11, color=BLUE, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    set_font(r2, size=10.5, color=NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_label_table(doc, rows, widths=(2700, 6660), header=None):
    table = doc.add_table(rows=1 if header else 0, cols=2)
    table.style = "Table Grid"
    if header:
        table.cell(0, 0).text = header[0]
        table.cell(0, 1).text = header[1]
        for cell in table.rows[0].cells:
            set_cell_shading(cell, LIGHT_BLUE)
            for p in cell.paragraphs:
                for r in p.runs:
                    set_font(r, size=10.5, color=BLUE, bold=True)
    for label, detail in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = detail
        set_cell_shading(cells[0], LIGHT_GRAY)
        for p in cells[0].paragraphs:
            for r in p.runs:
                set_font(r, size=10.5, color=BLUE, bold=True)
        for p in cells[1].paragraphs:
            for r in p.runs:
                set_font(r, size=10.5, color=NAVY)
    set_table_geometry(table, list(widths))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_matrix(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        set_cell_shading(cell, LIGHT_BLUE)
        for p in cell.paragraphs:
            for r in p.runs:
                set_font(r, size=10, color=BLUE, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            for p in cells[i].paragraphs:
                for r in p.runs:
                    set_font(r, size=9.8, color=NAVY)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


doc = Document()
style_doc(doc)
section = doc.sections[0]
section.header.paragraphs[0].text = "DEVORA Sentinel - Notice procedure client"
section.header.paragraphs[0].runs[0].font.size = Pt(9)
section.header.paragraphs[0].runs[0].font.color.rgb = MUTED
section.footer.paragraphs[0].text = "Document interne DEVORA - ne pas transmettre tel quel aux prospects"
section.footer.paragraphs[0].runs[0].font.size = Pt(9)
section.footer.paragraphs[0].runs[0].font.color.rgb = MUTED

title = doc.add_paragraph()
title.paragraph_format.space_before = Pt(18)
title.paragraph_format.space_after = Pt(4)
run = title.add_run("DEVORA Sentinel")
set_font(run, size=24, color=NAVY, bold=True)
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(18)
run = subtitle.add_run("Notice procedure client - Early Access, qualification, activation et support")
set_font(run, size=13, color=MUTED)

add_label_table(
    doc,
    [
        ("Version", "1.0 - preparation lancement"),
        ("Date", "21 mai 2026"),
        ("Usage", "Procedure interne pour traiter les demandes Early Access DEVORA Sentinel"),
        ("Principe", "Acces manuel, controle humain, donnees locales par defaut, aucune promesse agressive."),
    ],
    widths=(1900, 7460),
)

add_callout(
    doc,
    "Objectif de cette notice",
    "Avoir sous la main une procedure simple pour recevoir une demande, verifier si le client est eligible, lui repondre clairement, activer l'acces manuellement et assurer le suivi sans exposer les details internes du produit.",
    LIGHT_CYAN,
)

add_heading(doc, "1. Message clair a garder en tete", 1)
add_p(
    doc,
    "DEVORA Sentinel est un outil local premium qui aide a surveiller, visualiser et mieux comprendre certaines activites techniques, avec une logique defensive, lisible et controlee.",
)
add_label_table(
    doc,
    [
        ("Ce que c'est", "Un assistant local de surveillance defensive, de suivi d'incidents et de visibilite systeme."),
        ("Ce que ce n'est pas", "Ce n'est pas un spyware, pas un cloud invasif, pas une IA autonome qui agit sans validation."),
        ("Donnees", "Les informations suivies restent sur la machine par defaut. Les veilles externes doivent rester limitees aux sources autorisees."),
        ("Decision", "Les actions sensibles restent controlees par l'utilisateur et validees humainement."),
    ],
    widths=(2100, 7260),
)

add_heading(doc, "2. Parcours complet d'une demande", 1)
for step in [
    "Le visiteur arrive sur la page DEVORA Sentinel et lit la promesse locale, defensive et controlee.",
    "Il remplit le formulaire Early Access avec nom, email, OS, annee du PC, RAM, usage et option theme si besoin.",
    "La route /api/early-access/sentinel recoit la demande.",
    "Si AlexCore est configure, la demande est transmise a l'endpoint interne. Sinon, le site affiche l'email de secours.",
    "Tu relis la demande manuellement avant toute activation.",
    "Tu verifies la compatibilite, le contexte et l'absence d'information sensible dans le message.",
    "Tu reponds avec un email adapte : accepte, besoin d'informations, appel necessaire ou acces reporte.",
    "Tu notes la demande dans ton suivi interne et tu gardes une trace des activations.",
]:
    add_number(doc, step)

add_heading(doc, "3. Checklist de qualification", 1)
add_matrix(
    doc,
    ["Point a verifier", "Pourquoi c'est important", "Decision"],
    [
        ("Email et identite", "Pouvoir recontacter la personne et eviter les demandes fantaisistes.", "OK / a clarifier"),
        ("Machine", "Windows 10/11, PC x64 recent, RAM suffisante et usage coherent.", "OK / appel"),
        ("Usage", "Comprendre si c'est personnel, homelab, developpeur, TPE, agence ou equipe.", "Orienter vers pack"),
        ("Attentes", "Eviter les demandes de promesse antivirus totale ou d'automatisation agressive.", "Recadrer"),
        ("Confidentialite", "Ne pas recevoir de mots de passe, tokens, cles ou details sensibles.", "Supprimer/recadrer"),
    ],
    [2400, 4680, 2280],
)

add_callout(
    doc,
    "Regle de securite",
    "Si un prospect envoie un mot de passe, une cle API, un token, un fichier sensible ou une capture interne confidentielle, ne l'utilise pas. Demande une version expurgee et garde la discussion sur le besoin, pas sur les secrets.",
    "FFF7ED",
)

add_heading(doc, "4. Packs et options a annoncer", 1)
add_matrix(
    doc,
    ["Pack", "Prix", "Pour qui", "A inclure dans la reponse"],
    [
        ("Early Access", "0€", "50 premiers acces qualifies", "1 mois inclus, activation manuelle, support direct fondateur."),
        ("Solo", "14,90€/mois", "Independant, createur, homelab", "Vue incidents, historique, surveillance locale, 2 projets max."),
        ("Developer", "29€/mois", "Developpeur, studio, environnement technique", "Sentinel Watch, veille defensive, exports avances, 5 projets max."),
        ("Business", "79€/mois", "TPE, agence, equipe operationnelle", "Multi-projets, validation avancee, escalade incidents, support prioritaire."),
        ("Enterprise", "Sur devis", "Equipe avec besoin specifique", "Installation accompagnee, regles dediees, on-prem, integration specifique."),
    ],
    [1700, 1350, 2710, 3600],
)
add_label_table(
    doc,
    [
        ("Theme personnalise", "+29€"),
        ("Installation accompagnee", "+49€"),
        ("Configuration optimisation locale", "+29€"),
        ("Surveillance avancee supplementaire", "+9€/mois"),
        ("Historique incidents etendu", "+5€/mois"),
    ],
    widths=(4300, 5060),
    header=("Option", "Prix public"),
)

add_heading(doc, "5. Reponses types", 1)
add_heading(doc, "Demande acceptee Early Access", 2)
add_p(
    doc,
    "Bonjour [Prenom], merci pour votre demande. Votre profil semble coherent pour l'Early Access DEVORA Sentinel. L'acces est manuel et limite a 50 personnes, avec 1 mois inclus. Avant activation, je confirme la compatibilite de votre machine et le cadre d'utilisation. DEVORA Sentinel fonctionne principalement en local : les donnees suivies restent sur la machine par defaut, avec controle humain et sans collecte cachee.",
)
add_heading(doc, "Besoin d'informations avant activation", 2)
add_p(
    doc,
    "Bonjour [Prenom], merci pour votre demande. Avant de valider l'acces, il me manque quelques informations : systeme d'exploitation, annee approximative du PC, memoire RAM et usage principal. Le but est de vous orienter proprement et d'eviter une installation qui ne correspondrait pas a votre contexte.",
)
add_heading(doc, "Demande a recadrer", 2)
add_p(
    doc,
    "Bonjour [Prenom], DEVORA Sentinel est un outil de surveillance defensive et de visibilite locale. Il ne remplace pas un antivirus, ne promet pas une protection totale et n'agit pas automatiquement de maniere agressive. Si votre besoin est plutot un audit, une installation accompagnee ou une supervision avancee, je peux vous proposer un appel court pour cadrer le bon format.",
)

add_heading(doc, "6. Procedure technique avant mise en production", 1)
for item in [
    "Executer npm run build et corriger toute erreur avant publication.",
    "Verifier que la page /portfolio/devora-sentinel#early-access charge bien le formulaire.",
    "Configurer DEVORA_CONTACT_EMAIL et NEXT_PUBLIC_DEVORA_CONTACT_EMAIL.",
    "Configurer ALEXCORE_EARLY_ACCESS_URL et ALEXCORE_API_TOKEN uniquement si le relais interne est pret.",
    "Tester une demande avec un email a toi, puis verifier le comportement succes ou email de secours.",
    "Ne jamais publier d'endpoint interne, token, workflow Sentinel reel ou mecanique AlexCore.",
    "Apres mise en ligne, refaire un test mobile du formulaire et du bouton d'envoi.",
]:
    add_bullet(doc, item)

add_heading(doc, "7. Procedure d'activation client", 1)
for step in [
    "Creer ou reserver l'acces client dans ton suivi interne.",
    "Associer nom, email, pack, date de debut, date de fin Early Access et options choisies.",
    "Envoyer l'email d'acceptation avec le cadre local, les limites et les prochaines etapes.",
    "Planifier l'installation ou envoyer le lien quand le mode de distribution est pret.",
    "Demander un retour simple apres les premiers jours : installation, comprehension, incidents, performance.",
    "A la fin du mois inclus, proposer Solo, Developer, Business ou un devis selon usage reel.",
]:
    add_number(doc, step)

add_heading(doc, "8. Limites a ne pas promettre", 1)
add_matrix(
    doc,
    ["Ne pas dire", "Dire plutot"],
    [
        ("Sentinel bloque toutes les menaces.", "Sentinel aide a mieux voir, comprendre et traiter certains signaux defensifs."),
        ("L'IA gere tout automatiquement.", "L'IA aide a lire, resumer et prioriser, avec controle humain."),
        ("Aucune donnee ne sort jamais.", "Les donnees restent locales par defaut ; seules certaines veilles externes autorisees utilisent Internet."),
        ("C'est un antivirus.", "C'est un outil de visibilite et d'assistance defensive, complementaire aux protections existantes."),
        ("Demonstration complete publique.", "Les demonstrations detaillees sont reservees aux rendez-vous ou acces qualifies."),
    ],
    [3600, 5760],
)

add_heading(doc, "9. Suivi support", 1)
add_label_table(
    doc,
    [
        ("Early Access", "Support direct fondateur. Repondre vite, mais garder une trace des demandes et incidents."),
        ("Solo", "Support standard. Traiter les questions d'installation, comprehension et exports simples."),
        ("Developer", "Support plus technique. Verifier veilles, logs enrichis et projets suivis."),
        ("Business", "Support prioritaire. Traiter escalation, multi-projets et besoin d'equipe."),
        ("Enterprise", "Support personnalise. Passer par cadrage specifique et devis."),
    ],
    widths=(2200, 7160),
)

add_callout(
    doc,
    "Phrase finale a garder",
    "DEVORA Sentinel n'est pas une IA obscure dans le cloud. C'est un outil local, defensif et controle, pense pour rendre les signaux techniques plus lisibles et plus faciles a traiter.",
    LIGHT_CYAN,
)

doc.save(OUT)
print(OUT)
